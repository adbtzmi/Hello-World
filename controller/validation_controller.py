# -*- coding: utf-8 -*-
"""
controller/validation_controller.py
====================================
Validation Controller - Phase 3C

Bridges ValidationTab (View) <-> jira_analyzer.py + AutoConsolidator (Model).
Handles validation document generation, risk assessment, and checkout
result consolidation (trace analysis, validation doc, spool summary).
"""

import os
import glob
import logging
import threading

logger = logging.getLogger("bento_app")

CHECKOUT_RESULTS_FOLDER = r"P:\temp\BENTO\CHECKOUT_RESULTS"


class ValidationController:
    """
    Bridges ValidationTab <-> jira_analyzer.py + AutoConsolidator.

    When the user clicks "Generate Validation & Risk Assessment":
      1. Locate trace files from CHECKOUT_RESULTS (using JIRA key)
      2. Run AutoConsolidator (trace analysis → validation doc → spool
         summary → manifest → rule-based risk → AI validation → AI risk)
      3. Run JIRA-based risk assessment (existing workflow)
      4. Return unified results to the Validation tab
    """

    def __init__(self, context, workflow_ctrl, chat_ctrl):
        self.context = context
        self.workflow = workflow_ctrl
        self.chat = chat_ctrl
        self.analyzer = context.analyzer
        self._running = False
        logger.info("ValidationController initialised.")

    def is_running(self):
        return self._running

    # ──────────────────────────────────────────────────────────────────────
    # ASSESS RISKS  (unified: JIRA + trace consolidation)
    # ──────────────────────────────────────────────────────────────────────

    def assess_risks(self, issue_key, callback):
        """
        Generate validation document and risk assessment.

        Args:
            issue_key: JIRA issue key
            callback: Function to call with result dict
        """
        if self._running:
            logger.warning("ValidationController: already running")
            return

        self._running = True
        threading.Thread(
            target=self._assess_risks_thread,
            args=(issue_key, callback),
            daemon=True,
            name="bento-validation"
        ).start()

    def _assess_risks_thread(self, issue_key, callback):
        """Background thread for unified risk assessment."""
        try:
            self._log(f"[Generating Validation Document for {issue_key}]")

            # ── Phase A: Auto-consolidate checkout results ─────────────
            consolidation = self._run_checkout_consolidation(issue_key)

            # ── Phase B: JIRA-based analysis ───────────────────────────
            jira_result = self._run_jira_analysis(issue_key)

            # ── Merge results ──────────────────────────────────────────
            result = {
                'success': True,
                'consolidation': consolidation,
                'jira_assessment': jira_result,
            }

            # If JIRA analysis produced an assessment, include it
            if jira_result and jira_result.get('success'):
                result['assessment'] = jira_result.get('assessment', '')

            # If consolidation failed and JIRA also failed, report error
            no_consolidation = (not consolidation
                                or not consolidation.get('success'))
            no_jira = (not jira_result
                       or not jira_result.get('success'))
            if no_consolidation and no_jira:
                result['success'] = False
                errors = []
                if consolidation:
                    errors.extend(consolidation.get('errors', []))
                if jira_result:
                    errors.append(jira_result.get('error', ''))
                result['error'] = '; '.join(e for e in errors if e)

            self._callback(callback, result)

        except Exception as e:
            logger.error(f"Validation error: {e}")
            self._callback(callback, {
                'success': False,
                'error': str(e)
            })
        finally:
            self._running = False

    # ──────────────────────────────────────────────────────────────────────
    # PHASE A — Checkout result consolidation (trace-based)
    # ──────────────────────────────────────────────────────────────────────

    def _run_checkout_consolidation(self, issue_key):
        """
        Find trace files from CHECKOUT_RESULTS and run AutoConsolidator.

        Returns consolidation result dict or None if no trace files found.
        """
        try:
            from model.analyzers.auto_consolidator import AutoConsolidator
        except ImportError:
            self._log("⚠ AutoConsolidator not available — skipping")
            return None

        # Find trace files
        trace_info = self._find_checkout_trace_files(issue_key)
        if not trace_info:
            self._log("ℹ No checkout trace files found — "
                      "skipping trace consolidation")
            return None

        collection_dir = trace_info['collection_dir']
        trace_file = trace_info['trace_file']
        database_path = trace_info.get('database_path')
        mid = trace_info.get('mid', '')

        self._log(f"📊 Found trace file: {os.path.basename(trace_file)}")
        self._log(f"   Collection dir: {collection_dir}")

        # Resolve AI client if available
        ai_client = None
        try:
            controller = self.context.controller
            if hasattr(controller, '_ai_client'):
                ai_client = controller._ai_client
            elif hasattr(self.context, 'ai_client'):
                ai_client = self.context.ai_client
        except Exception:
            pass

        # Get workflow context for AI analysis
        jira_context = None
        test_scenarios = None
        code_changes = None
        impact_analysis = None
        try:
            jira_text = self.workflow.get_workflow_step("JIRA_ANALYSIS")
            if jira_text:
                jira_context = {"analysis": jira_text, "key": issue_key}
            test_scenarios = self.workflow.get_workflow_step("TEST_SCENARIOS")
            impact_analysis = self.workflow.get_workflow_step("IMPACT_ANALYSIS")
        except Exception:
            pass

        # Run consolidation
        self._log("📊 Running auto-consolidation...")
        consolidator = AutoConsolidator(
            template_path="template_validation.docx",
            ai_client=ai_client,
        )

        results = consolidator.consolidate(
            collection_dir=collection_dir,
            trace_file_path=trace_file,
            database_path=database_path,
            jira_key=issue_key,
            mid=mid,
            generate_validation_doc=True,
            generate_spool_summary=True,
            generate_manifest=True,
            perform_risk_assessment=True,
            perform_ai_validation=ai_client is not None,
            perform_ai_risk_assessment=ai_client is not None,
            jira_context=jira_context,
            test_scenarios=test_scenarios,
            code_changes=code_changes,
            impact_analysis=impact_analysis,
            log_callback=lambda msg: self._log(msg),
        )

        if results.get('success'):
            self._log("✓ Auto-consolidation completed successfully")
        else:
            errors = results.get('errors', [])
            self._log(f"⚠ Auto-consolidation had errors: "
                      f"{', '.join(errors)}")

        return results

    def _find_checkout_trace_files(self, issue_key):
        """
        Search CHECKOUT_RESULTS for trace files matching the JIRA key.

        Scans P:\\temp\\BENTO\\CHECKOUT_RESULTS\\*\\<issue_key>\\
        for DispatcherDebug*.txt or *TraceFile*.txt files.

        Returns dict with 'collection_dir', 'trace_file', 'database_path',
        'mid' or None if not found.
        """
        if not os.path.isdir(CHECKOUT_RESULTS_FOLDER):
            return None

        # Search all tester folders for this JIRA key
        best_trace = None
        best_mtime = 0

        try:
            for tester_dir in os.listdir(CHECKOUT_RESULTS_FOLDER):
                tester_path = os.path.join(
                    CHECKOUT_RESULTS_FOLDER, tester_dir)
                if not os.path.isdir(tester_path):
                    continue

                jira_path = os.path.join(tester_path, issue_key)
                if not os.path.isdir(jira_path):
                    continue

                # Search recursively for trace files
                for root, dirs, files in os.walk(jira_path):
                    for fname in files:
                        is_trace = (
                            fname.startswith("DispatcherDebug")
                            and fname.endswith(".txt")
                        ) or "TraceFile" in fname
                        if not is_trace:
                            continue

                        full_path = os.path.join(root, fname)
                        try:
                            mtime = os.path.getmtime(full_path)
                        except OSError:
                            mtime = 0

                        if mtime > best_mtime:
                            best_mtime = mtime
                            best_trace = full_path
        except Exception as e:
            logger.warning(f"Error scanning CHECKOUT_RESULTS: {e}")
            return None

        if not best_trace:
            return None

        collection_dir = os.path.dirname(best_trace)

        # Find database file in same directory
        database_path = None
        for f in glob.glob(os.path.join(collection_dir, "*.db")):
            if "resultsManager" in os.path.basename(f):
                database_path = f
                break

        # Try to extract MID from directory name
        mid = os.path.basename(collection_dir)

        return {
            'collection_dir': collection_dir,
            'trace_file': best_trace,
            'database_path': database_path,
            'mid': mid,
        }

    # ──────────────────────────────────────────────────────────────────────
    # PHASE B — JIRA-based analysis (existing workflow)
    # ──────────────────────────────────────────────────────────────────────

    def _run_jira_analysis(self, issue_key):
        """
        Run JIRA-based risk assessment (existing workflow).

        Returns dict with 'success', 'assessment', 'template_file'
        or dict with 'success': False, 'error'.
        """
        try:
            # Initialize workflow
            self.workflow.init_workflow_file(issue_key)

            # Get repository path
            repo_path = self.workflow.get_workflow_step("REPOSITORY_PATH")
            if not repo_path:
                self._log("ℹ Repository path not in workflow — "
                          "skipping JIRA-based analysis")
                return {'success': False,
                        'error': "Repository path not found in workflow"}

            # Check for template
            template_file = "template_validation.docx"
            if os.path.exists(template_file):
                self._log(f"✓ Found template: {template_file}")
                self._populate_template(issue_key, template_file, repo_path)

            # Index repository
            self._log("Indexing repository...")
            repo_index = self.analyzer.index_repository(repo_path)

            # Get JIRA analysis
            jira_analysis_text = self.workflow.get_workflow_step(
                "JIRA_ANALYSIS")
            if jira_analysis_text:
                self._log("✓ Using existing JIRA analysis from workflow")
                jira_analysis = {
                    'success': True, 'analysis': jira_analysis_text}
            else:
                self._log("⚠ JIRA analysis not found, fetching...")
                issue_data = self.analyzer.fetch_jira_issue(issue_key)
                if not issue_data:
                    return {'success': False,
                            'error': f"Failed to fetch JIRA issue {issue_key}"}

                jira_analysis = self.analyzer.analyze_jira_request(issue_data)
                if not jira_analysis.get('success'):
                    return {
                        'success': False,
                        'error': ("Failed to analyze JIRA: "
                                  + jira_analysis.get('error', ''))
                    }

                self.workflow.save_workflow_step(
                    "JIRA_ANALYSIS", jira_analysis['analysis'])

            # Get impact analysis
            impact_analysis_text = self.workflow.get_workflow_step(
                "IMPACT_ANALYSIS")
            if impact_analysis_text:
                self._log("✓ Using existing impact analysis from workflow")
                impact_analysis = {
                    'success': True,
                    'impact_analysis': impact_analysis_text}
            else:
                self._log("⚠ Impact analysis not found, generating...")
                impact_analysis = self.analyzer.analyze_code_impact(
                    repo_path, repo_index, jira_analysis
                )
                if not impact_analysis.get('success'):
                    return {
                        'success': False,
                        'error': ("Failed to analyze impact: "
                                  + impact_analysis.get('error', ''))
                    }

                self.workflow.save_workflow_step(
                    "IMPACT_ANALYSIS",
                    impact_analysis['impact_analysis'])

            # Generate risk assessment
            self._log("Generating risk assessment with AI...")
            risk_assessment = self.analyzer.assess_risks(
                jira_analysis, impact_analysis)

            if risk_assessment.get('success'):
                assessment_text = risk_assessment.get('risk_assessment', '')
                self._log("✓ Risk assessment generated")
                return {
                    'success': True,
                    'assessment': assessment_text,
                    'template_file': f"Validation_{issue_key}.docx",
                }
            else:
                return {
                    'success': False,
                    'error': ("Risk assessment failed: "
                              + risk_assessment.get('error', ''))
                }

        except Exception as e:
            logger.error(f"JIRA analysis error: {e}")
            return {'success': False, 'error': str(e)}

    def _populate_template(self, issue_key, template_file, repo_path):
        """Populate validation template with workflow data."""
        try:
            from docx import Document

            doc = Document(template_file)

            # Get workflow data
            jira_data = (self.workflow.get_workflow_step("JIRA_ISSUE_DATA")
                         or "N/A")
            jira_analysis = (self.workflow.get_workflow_step("JIRA_ANALYSIS")
                             or "N/A")
            impact_analysis = (
                self.workflow.get_workflow_step("IMPACT_ANALYSIS") or "N/A")
            test_scenarios = (
                self.workflow.get_workflow_step("TEST_SCENARIOS") or "N/A")

            # Replace placeholders
            replacements = {
                '{{ISSUE_KEY}}': issue_key,
                '{{JIRA_DATA}}': jira_data[:500],
                '{{JIRA_ANALYSIS}}': jira_analysis[:500],
                '{{IMPACT_ANALYSIS}}': impact_analysis[:500],
                '{{TEST_SCENARIOS}}': test_scenarios[:500],
                '{{REPO_PATH}}': repo_path,
            }

            for paragraph in doc.paragraphs:
                for key, value in replacements.items():
                    if key in paragraph.text:
                        paragraph.text = paragraph.text.replace(key, value)

            # Save populated template
            output_file = f"Validation_{issue_key}.docx"
            doc.save(output_file)
            self._log(f"✓ Validation document saved: {output_file}")

        except ImportError:
            self._log("⚠ python-docx not installed, skipping template")
        except Exception as e:
            logger.error(f"Template population error: {e}")
            self._log(f"⚠ Template population failed: {e}")

    # ──────────────────────────────────────────────────────────────────────
    # FINALIZE ASSESSMENT
    # ──────────────────────────────────────────────────────────────────────

    def finalize_assessment(self, issue_key, assessment_text):
        """Save finalized risk assessment to workflow."""
        try:
            self.workflow.save_workflow_step("RISK_ASSESSMENT",
                                            assessment_text)
            self._log(f"✓ Risk assessment saved to workflow")
            return True
        except Exception as e:
            logger.error(f"Failed to finalize assessment: {e}")
            return False

    # ──────────────────────────────────────────────────────────────────────
    # HELPERS
    # ──────────────────────────────────────────────────────────────────────

    def _log(self, message):
        """Thread-safe logging."""
        if self.context.log_callback:
            self.context.root.after(
                0, lambda m=message: self.context.log_callback(m))

    def _callback(self, callback, result):
        """Thread-safe callback."""
        if callback:
            self.context.root.after(0, lambda r=result: callback(r))
